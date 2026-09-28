"""
app.py (Mimir)
Chat UI for Mimir, the local Wiki LLM -- retrieval-augmented Q&A over
heimdall's README + module docstrings, running entirely through a local
Ollama instance (localhost:11434 only, nothing leaves this machine).

Named after the Norse keeper of wisdom, whose counsel could always be
sought -- the point of this app: institutional "why" knowledge that stays
answerable, not buried in scattered code comments and one person's memory.

Deliberately calls the index + Ollama directly in-process rather than
through a separate API service -- this is a single-user local tool, so a
second HTTP hop would only add moving parts without solving a real problem.

Run via mimir/run.bat, or manually:
    venv/Scripts/python.exe -m streamlit run app.py --server.port 8600
"""

import base64
import io
import json
import time

import chromadb
import ollama
import openpyxl
import streamlit as st

import bench
import indexer
import memory_store
import tools

st.set_page_config(page_title="Mimir - SDA", layout="wide")

# qwen3:8b replaced qwen2.5:7b-instruct and hermes3:8b (both removed) after a
# 3-run comparison: matched or beat qwen2.5 on every axis (see memory of this
# comparison). deepseek-r1:8b was also tested and removed -- it never called
# the tool for direct data requests (0/3), instead fabricating an empty
# template mimicking the tool's own output labels with no real data behind it.
MODELS = ["qwen3:8b"]
TOP_K = 4
CATAT_PREFIX = "catat:"
MAX_TOOL_ROUNDS = 2  # a lookup then a possible follow-up (e.g. search name -> fetch outlet)
MAX_HISTORY_MESSAGES = 20  # ~10 exchanges -- plenty of headroom left in either
                           # model's context window, just bounded so a very
                           # long session doesn't keep growing the prompt forever

SYSTEM_PROMPT = """Nama Anda adalah Mimir -- asisten wiki internal (Wiki LLM) milik \
PT SDA yang berjalan sepenuhnya lokal di komputer ini. Anda BUKAN Heimdall -- \
Heimdall adalah nama aplikasi web TERPISAH (folder heimdall/) yang mengolah data \
OMSHAR/EAO, dan tugas Anda adalah menjawab pertanyaan TENTANG Heimdall berdasarkan \
dokumentasinya, bukan menjadi Heimdall itu sendiri. Kalau ditanya siapa Anda, \
jawab: Anda adalah Mimir.

Untuk pertanyaan tentang CARA KERJA sistem, jawab HANYA berdasarkan referensi \
dokumentasi yang diberikan -- jangan mengarang informasi yang tidak ada di situ. \
Referensi bisa berasal dari README/docstring kode Heimdall, ATAU dari memory yang \
sudah dicatat sebelumnya (keputusan, koreksi, proses, atau struktur file Excel) -- \
keduanya sama validnya. Untuk permintaan DATA outlet, referensi dokumentasi tidak \
relevan: langsung panggil tool, jangan menjelaskan langkah-langkah cara memakai \
Heimdall atau menulis kode.

PENTING soal angka: Anda punya dua tool untuk data outlet nyata dari Heimdall -- \
cari_outlet (data omset satu outlet berdasarkan site number) dan cari_nama_outlet \
(cari site number dari nama outlet). Kalau user meminta data/angka outlet, PANGGIL tool, \
jangan menebak. Semua angka yang Anda sebut HARUS persis dari hasil tool -- jangan \
mengarang, membulatkan, atau menghitung sendiri. Kalau tool bilang tidak ditemukan atau \
gagal, sampaikan apa adanya. Kalau user menyebut nama outlet tanpa site number, panggil \
cari_nama_outlet dulu. Anda TIDAK punya akses ke klaim SKU atau data Admin lainnya -- \
untuk itu arahkan user ke halaman terkait di Heimdall. Untuk pertanyaan tentang cara \
kerja sistem (bukan data outlet), jawab dari dokumentasi tanpa memanggil tool.

Aturan tool: (1) JANGAN berjanji akan mengambil data ("tunggu sebentar", "saya akan \
mencoba") -- langsung panggil tool di giliran yang sama. (2) Untuk pertanyaan LANJUTAN \
tentang angka (mis. "bulan terakhir", "brand apa saja"), panggil tool lagi -- jangan \
membaca angka dari jawaban Anda sebelumnya di percakapan. (3) Sampaikan hasil tool apa \
adanya dan ringkas; kalau outlet tidak ditemukan, katakan langsung "Site ... tidak \
ditemukan". Gunakan baris "BULAN TERAKHIR" dan "Nilai per brand di bulan terakhir" dari \
hasil tool untuk pertanyaan soal bulan terakhir.

Jawab dalam Bahasa Indonesia, singkat dan jelas."""

MEMORY_DRAFT_PROMPT = """Anda membantu menyusun catatan memory dari pesan user. \
Baca pesan user di bawah dan hasilkan draft memory dalam format JSON persis sesuai \
schema yang diberikan. Field "type" HARUS salah satu dari: "decision", "correction", \
"process". Field "title" singkat (maks 8 kata). Field "tags" 1-4 kata kunci relevan. \
Field "content" adalah isi memory yang rapi (boleh sedikit dirapikan dari pesan asli, \
tapi JANGAN menambahkan informasi yang tidak disebutkan user)."""

MEMORY_DRAFT_SCHEMA = {
    "type": "object",
    "properties": {
        "type": {"type": "string", "enum": ["decision", "correction", "process"]},
        "title": {"type": "string"},
        "tags": {"type": "array", "items": {"type": "string"}},
        "content": {"type": "string"},
    },
    "required": ["type", "title", "tags", "content"],
}


@st.cache_resource
def get_chroma_collection():
    client = chromadb.PersistentClient(path=str(indexer.INDEX_DIR))
    return client.get_collection(indexer.COLLECTION_NAME)


def retrieve_context(query: str, top_k: int = TOP_K) -> list[dict]:
    collection = get_chroma_collection()
    query_embedding = ollama.embed(model=indexer.EMBED_MODEL, input=query)["embeddings"][0]
    results = collection.query(query_embeddings=[query_embedding], n_results=top_k)
    return [
        {"text": doc, "source": meta["source"]}
        for doc, meta in zip(results["documents"][0], results["metadatas"][0])
    ]


def build_prompt(query: str, chunks: list[dict]) -> str:
    context = "\n\n".join(f"[{c['source']}]\n{c['text']}" for c in chunks)
    # Question first, docs second and labelled as conditional: with the docs
    # first, retrieved README text about "how to search an outlet in Heimdall's
    # UI" pulled the model into explaining those steps instead of calling the
    # tool for a data request.
    return (
        f"Pertanyaan: {query}\n\n"
        "Referensi dokumentasi (hanya untuk pertanyaan tentang cara kerja sistem; "
        f"untuk permintaan data outlet abaikan dan panggil tool):\n\n{context}"
    )


def build_history(exclude_last: bool = True) -> list[dict]:
    """Prior turns as plain {role, content} dicts for ollama.chat(), capped
    to MAX_HISTORY_MESSAGES. Without this, every question was previously
    sent as an isolated conversation with no memory of anything said
    before it in the same session -- the chat log LOOKED continuous
    (Streamlit just replays session_state.messages on screen) but the
    model itself never saw any of it, so a follow-up like "kenapa kamu
    bilang begitu" (why did you say that) had literally nothing to refer
    back to and just repeated a generic answer.

    exclude_last=True drops the just-appended current question, since
    that gets sent separately (with fresh retrieved context attached),
    not replayed as bare history."""
    msgs = st.session_state.messages[:-1] if exclude_last else st.session_state.messages
    history = [{"role": m["role"], "content": m["content"]} for m in msgs]
    return history[-MAX_HISTORY_MESSAGES:]


def draft_memory(model: str, raw_text: str) -> tuple[dict, dict]:
    """Sends the user's `catat:` message to the LLM asking for structured
    JSON output (type/title/tags/content) -- a fixed schema is far more
    reliable than parsing free text out of an 8B model's response.

    Returns (draft, timing) -- timing logged via bench.log() same as the
    main chat flow, so `catat:` drafts count toward the model benchmark too
    (kind="catat_draft", separate from "chat" so sidebar averages aren't
    skewed by a different kind of call)."""
    t0 = time.perf_counter()
    resp = ollama.chat(
        model=model,
        messages=[
            {"role": "system", "content": MEMORY_DRAFT_PROMPT},
            {"role": "user", "content": raw_text},
        ],
        format=MEMORY_DRAFT_SCHEMA,
    )
    elapsed = time.perf_counter() - t0
    timing = bench.log(
        "catat_draft", model, retrieval_s=0.0, generate_s=elapsed, total_s=elapsed,
        ollama_stats=resp if isinstance(resp, dict) else None,
    )
    return json.loads(resp["message"]["content"]), timing


def extract_excel_structure(file_bytes: bytes, filename: str) -> str:
    """Sheet names + header row per sheet -- structure only, never actual
    data rows, so nothing here can leak real figures into the index."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=False)
    lines = [f"File: {filename}", f"Sheets: {', '.join(wb.sheetnames)}", ""]
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        try:
            header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        except StopIteration:
            header_row = ()
        headers = [str(h) for h in header_row if h is not None]
        lines.append(f"## Sheet: {sheet_name}")
        lines.append(f"Columns: {', '.join(headers) if headers else '(kosong/tidak terbaca)'}")
        lines.append("")
    wb.close()
    return "\n".join(lines)


def render_copy_print_widget(png_bytes: bytes) -> None:
    """Copy-to-clipboard + Print buttons for a report PNG -- ported from
    heimdall/omset_search_app.py's Omset Seeker page rather than
    reimplemented, since that JS already went through real
    browser-compatibility debugging (clipboard API is blocked outside a
    "secure context" -- https:// or http://localhost -- with a
    contenteditable+execCommand fallback for LAN access; Print navigates a
    real <a> to a Blob URL instead of window.open(), since the latter is
    popup-blocked from an iframe). Each st.iframe() call runs in its own
    isolated iframe, so the same fixed element ids/function name are safe
    to reuse across multiple outlet lookups shown in one chat session.

    Uses st.iframe() (raw-HTML-string form), not components.html() -- the
    latter is deprecated in this Streamlit version. height="content" was
    tried first per st.iframe's own docs ("auto-sizes to content height")
    but left a large blank gap in practice for this snippet-style content
    (buttons + a hidden print stylesheet, not a full laid-out page) --
    reverted to a fixed height=50, same as the working components.html()
    call this replaced."""
    b64 = base64.b64encode(png_bytes).decode()
    st.iframe(
        f"""
        <style>
          .action-btn {{
            display: inline-block; padding: 0.5rem 1rem; font-size: 1rem; cursor: pointer;
            border-radius: 0.5rem; border: 1px solid #999; margin-right: 0.5rem;
            color: inherit; text-decoration: none; background: #f0f2f6;
            font-family: inherit;
          }}
        </style>
        <button class="action-btn" onclick="copyReportImage()">Copy to Clipboard</button>
        <a id="print-link" class="action-btn" href="#" target="_blank" rel="noopener noreferrer">Print</a>
        <span id="action-status" style="margin-left:0.5rem;"></span>
        <script>
        const reportImgSrc = "data:image/png;base64,{b64}";

        async function copyReportImage() {{
            const status = document.getElementById('action-status');
            status.textContent = 'Menyalin...';
            try {{
                const resp = await fetch(reportImgSrc);
                const blob = await resp.blob();
                if (!(window.isSecureContext && navigator.clipboard && navigator.clipboard.write)) {{
                    throw new Error('clipboard-api-unavailable');
                }}
                await navigator.clipboard.write([new ClipboardItem({{'image/png': blob}})]);
                status.textContent = 'Tersalin!';
                return;
            }} catch (e) {{
                // lanjut ke fallback di bawah
            }}
            try {{
                const ok = await copyImageLegacyFallback();
                status.textContent = ok ? 'Tersalin!' : 'Gagal copy -- coba Print atau klik kanan gambar > Copy image.';
            }} catch (e2) {{
                status.textContent = 'Gagal copy -- klik kanan gambar (kalau tabel HTML terlihat) atau pakai Print.';
            }}
        }}

        function copyImageLegacyFallback() {{
            return new Promise((resolve, reject) => {{
                const container = document.createElement('div');
                container.contentEditable = 'true';
                container.style.position = 'fixed';
                container.style.left = '-9999px';
                const img = document.createElement('img');
                img.onload = () => {{
                    document.body.appendChild(container);
                    container.appendChild(img);
                    const range = document.createRange();
                    range.selectNode(img);
                    const sel = window.getSelection();
                    sel.removeAllRanges();
                    sel.addRange(range);
                    let ok = false;
                    try {{
                        ok = document.execCommand('copy');
                    }} catch (e3) {{
                        ok = false;
                    }}
                    sel.removeAllRanges();
                    document.body.removeChild(container);
                    resolve(ok);
                }};
                img.onerror = () => reject(new Error('image-load-failed'));
                img.src = reportImgSrc;
            }});
        }}

        const printHtml =
            '<html><head><title>Print Laporan</title>' +
            '<style>' +
            '@page {{ size: landscape; margin: 5mm; }}' +
            'html, body {{ margin:0; padding:0; height:100%; }}' +
            '.print-page {{ width:287mm; height:200mm; display:flex; align-items:center; justify-content:center; }}' +
            '.print-page img {{ max-width:100%; max-height:100%; object-fit:contain; }}' +
            '</style>' +
            '</head>' +
            '<body>' +
            '<div class="print-page"><img src="' + reportImgSrc + '" onload="window.print()"></div>' +
            '</body></html>';
        const printBlob = new Blob([printHtml], {{type: 'text/html'}});
        document.getElementById('print-link').href = URL.createObjectURL(printBlob);
        </script>
        """,
        height=50,
    )


# ── Sidebar ──────────────────────────────────────────────────────────────
with st.sidebar:
    st.header("Mimir")
    st.caption("Lokal & offline -- tidak ada data yang keluar dari mesin ini.")

    model = st.selectbox("Model", MODELS, index=0)

    try:
        chunk_count = get_chroma_collection().count()
        st.caption(f"Index: {chunk_count} chunks")
    except Exception:
        st.warning("Index belum ada -- klik 'Bangun ulang index' di bawah.")

    if st.button("🔄 Bangun ulang index"):
        with st.spinner("Mengindeks ulang README + docstrings..."):
            get_chroma_collection.clear()
            n = indexer.build_index()
        st.success(f"Selesai -- {n} chunks terindeks.")
        st.rerun()

    if st.button("🗑️ Hapus riwayat chat"):
        st.session_state.messages = []
        st.rerun()

    with st.expander("📊 Benchmark model"):
        summary = bench.summarize_by_model()
        if not summary:
            st.caption("Belum ada respons chat yang tercatat.")
        else:
            for row in summary:
                tok = f", {row['avg_tokens_per_sec']} tok/s" if row["avg_tokens_per_sec"] else ""
                st.caption(f"**{row['model']}** -- {row['avg_total_s']}s rata-rata ({row['n']}x{tok})")

    st.divider()
    with st.expander("📊 Tambah Artifact Excel"):
        st.caption("Struktur (sheet/kolom) diambil otomatis. Tujuan file HARUS Anda tulis sendiri -- tidak pernah ditebak.")
        uploaded = st.file_uploader("File Excel", type=["xlsx"], key="excel_uploader")
        purpose = st.text_area("Tujuan file ini (opsional, tapi disarankan)", key="excel_purpose")
        if st.button("Simpan Artifact", disabled=uploaded is None):
            structure_text = extract_excel_structure(uploaded.getvalue(), uploaded.name)
            struct_path = memory_store.save_memory(
                "artifact-structure", uploaded.name, structure_text, source_file=uploaded.name,
            )
            indexer.add_single_chunk(structure_text, str(struct_path.relative_to(indexer.REPO_ROOT).as_posix()), "memory", "artifact-structure")
            saved_msg = f"Struktur tersimpan: `{struct_path.name}`"

            if purpose.strip():
                purpose_path = memory_store.save_memory(
                    "artifact-purpose", uploaded.name, purpose.strip(), source_file=uploaded.name,
                )
                indexer.add_single_chunk(purpose.strip(), str(purpose_path.relative_to(indexer.REPO_ROOT).as_posix()), "memory", "artifact-purpose")
                saved_msg += f" + tujuan tersimpan: `{purpose_path.name}`"
            else:
                saved_msg += " (tujuan belum diisi -- tidak dicatat, bisa ditambahkan lagi nanti dengan upload ulang)"

            get_chroma_collection.clear()
            st.success(saved_msg)

# ── Chat state ───────────────────────────────────────────────────────────
if "messages" not in st.session_state:
    st.session_state.messages = []
if "pending_memory" not in st.session_state:
    st.session_state.pending_memory = None
if "pending_memory_timing" not in st.session_state:
    st.session_state.pending_memory_timing = None

st.title("Mimir -- SDA Internal Wiki")
st.caption(
    "Tanya apa saja tentang cara kerja pipeline OMSHAR/EAO. Jawaban diambil dari "
    "README, komentar kode, dan memory yang sudah dicatat -- bukan dikarang. "
    f"Ketik `{CATAT_PREFIX} ...` untuk mencatat keputusan/koreksi/proses baru."
)

for msg_idx, msg in enumerate(st.session_state.messages):
    with st.chat_message(msg["role"]):
        st.markdown(msg["content"])
        for out_idx, out in enumerate(msg.get("tool_outputs") or []):
            with st.expander(f"🔧 {out['title']}", expanded=True):
                if out.get("html"):
                    st.markdown(out["html"], unsafe_allow_html=True)
                elif out["table"] is not None:
                    st.dataframe(out["table"])
                else:
                    st.caption(out["model_text"])
                if out.get("png_bytes"):
                    render_copy_print_widget(out["png_bytes"])
                    st.download_button(
                        "Download Gambar", out["png_bytes"], file_name=out.get("png_name", "laporan.png"),
                        mime="image/png", key=f"dl-hist-{msg_idx}-{out_idx}",
                    )
        if msg.get("timing"):
            st.caption(bench.format_caption(msg["timing"]))
        if msg.get("sources"):
            with st.expander("Sumber"):
                for s in msg["sources"]:
                    st.caption(f"📄 {s}")

# ── Pending memory confirmation (survives reruns until Simpan/Batal) ─────
if st.session_state.pending_memory:
    draft = st.session_state.pending_memory
    with st.chat_message("assistant"):
        st.markdown(f"**Draft memory -- {draft['type']}**")
        st.markdown(f"**{draft['title']}**")
        st.caption("Tags: " + ", ".join(draft["tags"]))
        st.markdown(draft["content"])
        if st.session_state.pending_memory_timing:
            st.caption(bench.format_caption(st.session_state.pending_memory_timing))
        col1, col2 = st.columns(2)
        if col1.button("✅ Simpan", key="save_memory"):
            path = memory_store.save_memory(draft["type"], draft["title"], draft["content"], tags=draft["tags"])
            rel = path.relative_to(indexer.REPO_ROOT).as_posix()
            indexer.add_single_chunk(draft["content"], rel, "memory", draft["type"])
            get_chroma_collection.clear()
            st.session_state.messages.append({
                "role": "assistant",
                "content": f"✅ Tersimpan sebagai `{path.name}` dan langsung bisa dicari.",
            })
            st.session_state.pending_memory = None
            st.session_state.pending_memory_timing = None
            st.rerun()
        if col2.button("❌ Batal", key="cancel_memory"):
            st.session_state.messages.append({"role": "assistant", "content": "❌ Dibatalkan, tidak disimpan."})
            st.session_state.pending_memory = None
            st.session_state.pending_memory_timing = None
            st.rerun()

# ── Handle new question ──────────────────────────────────────────────────
if question := st.chat_input("Tanya sesuatu, atau `catat: ...` untuk mencatat memory baru..."):
    st.session_state.messages.append({"role": "user", "content": question})
    with st.chat_message("user"):
        st.markdown(question)

    if question.strip().lower().startswith(CATAT_PREFIX):
        raw_text = question.strip()[len(CATAT_PREFIX):].strip()
        with st.chat_message("assistant"):
            with st.spinner("Menyusun draft memory..."):
                try:
                    draft, draft_timing = draft_memory(model, raw_text)
                    st.session_state.pending_memory = draft
                    st.session_state.pending_memory_timing = draft_timing
                except (json.JSONDecodeError, KeyError) as e:
                    st.error(f"Gagal menyusun draft -- coba tulis ulang lebih jelas. ({e})")
                except Exception as e:
                    st.error(
                        f"Gagal menghubungi Ollama ({type(e).__name__}: {e}) -- "
                        "pastikan Ollama sudah jalan (`ollama serve`) dan model-nya sudah di-pull."
                    )
        st.rerun()
    else:
        with st.chat_message("assistant"):
            t0 = time.perf_counter()
            with st.spinner("Mencari dokumentasi relevan..."):
                try:
                    chunks = retrieve_context(question)
                except Exception as e:
                    st.error(
                        f"Gagal menghubungi Ollama/index ({type(e).__name__}: {e}) -- "
                        "pastikan Ollama sudah jalan (`ollama serve`) dan index sudah dibangun."
                    )
                    st.stop()
            t_retrieved = time.perf_counter()
            sources = sorted({c["source"] for c in chunks})

            prompt = build_prompt(question, chunks)
            convo = [
                {"role": "system", "content": SYSTEM_PROMPT},
                *build_history(),
                {"role": "user", "content": prompt},
            ]

            # Simpan chunk TERAKHIR (done=True) -- itu yang bawa statistik asli
            # dari Ollama sendiri (eval_count/eval_duration, dst), dipakai
            # bench.log() untuk hitung tokens/sec yang akurat (bukan cuma
            # wall-clock Python yang ikut kehitung overhead Streamlit).
            final_part = {}
            tool_outputs = []  # tabel mentah hasil tool, ditampilkan langsung ke user

            def token_stream():
                """Streams the answer; if the model asks for a tool, runs it
                (Heimdall's real function -- see tools.py), feeds the result
                back, and streams the follow-up. Bounded by MAX_TOOL_ROUNDS so
                a model that keeps calling tools can't loop forever."""
                messages = convo
                for round_no in range(MAX_TOOL_ROUNDS + 1):
                    stream = ollama.chat(
                        model=model,
                        messages=messages,
                        # Last round offers no tools -> forces a plain-text answer.
                        tools=tools.TOOL_SCHEMAS if round_no < MAX_TOOL_ROUNDS else None,
                        stream=True,
                    )
                    text, calls = "", []
                    for part in stream:
                        if part.get("done"):
                            final_part.update(part)
                        msg = part.get("message", {})
                        content = msg.get("content", "")
                        if content:
                            text += content
                            yield content
                        calls.extend(msg.get("tool_calls") or [])
                    if not calls:
                        return
                    messages = [*messages, {"role": "assistant", "content": text, "tool_calls": calls}]
                    for call in calls:
                        result = tools.run_tool(call.function.name, dict(call.function.arguments or {}))
                        tool_outputs.append(result)
                        messages.append({"role": "tool", "tool_name": call.function.name, "content": result["model_text"]})

            try:
                answer = st.write_stream(token_stream())
            except Exception as e:
                st.error(
                    f"Gagal menghubungi Ollama ({type(e).__name__}: {e}) -- "
                    "pastikan Ollama sudah jalan (`ollama serve`) dan model-nya sudah di-pull."
                )
                st.stop()
            t_end = time.perf_counter()

            for out_idx, out in enumerate(tool_outputs):
                with st.expander(f"🔧 {out['title']}", expanded=True):
                    if out.get("html"):
                        st.markdown(out["html"], unsafe_allow_html=True)
                    elif out["table"] is not None:
                        st.dataframe(out["table"])
                    else:
                        st.caption(out["model_text"])
                    if out.get("png_bytes"):
                        render_copy_print_widget(out["png_bytes"])
                        st.download_button(
                            "Download Gambar", out["png_bytes"], file_name=out.get("png_name", "laporan.png"),
                            mime="image/png", key=f"dl-live-{out_idx}",
                        )

            timing = bench.log(
                "chat", model,
                retrieval_s=t_retrieved - t0,
                generate_s=t_end - t_retrieved,
                total_s=t_end - t0,
                ollama_stats=final_part or None,
            )
            st.caption(bench.format_caption(timing))

            with st.expander("Sumber"):
                for s in sources:
                    st.caption(f"📄 {s}")

        st.session_state.messages.append({
            "role": "assistant", "content": answer, "sources": sources, "timing": timing,
            "tool_outputs": tool_outputs,
        })
